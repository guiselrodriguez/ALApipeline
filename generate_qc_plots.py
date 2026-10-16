"""
Standardized QA/QC plot generator -- reads a house's H<N>.csv (output of
process_hbrl_data.py's `sensors` command) plus the House_charac file, and
produces the 4 fixed report figures the team asked for. Unlike test_plot.py
(which is for picking any variable combo you want to look at), this always
makes the same 4 figures, every time, for QA/QC.

Every figure:
 - shares one time axis across its panels (shown in PDT, local time)
 - has a red dashed line marking the exact moment the house got the stove
   (House_charac's Visit 2 timestamp, converted to PDT)
 - labels "pre" and "post" on the bottom panel
 - titled with the house number

Usage:
    python generate_qc_plots.py <H<N>.csv> <House_charac.xlsx> --output-dir qc_plots
"""

import argparse
import os
from datetime import timedelta

import openpyxl
import pandas as pd
import matplotlib.pyplot as plt

from process_hbrl_data import find_house_charac_header_row, parse_visit_value

PDT_OFFSET = timedelta(hours=7)


def get_stove_install_time(house_charac_file, house_id):
    """Visit 2's exact timestamp, converted to PDT -- this is literally
    when the house received the induction stove."""
    wb = openpyxl.load_workbook(house_charac_file, data_only=True)
    ws = wb["Sheet1"]
    header_row_number = find_house_charac_header_row(ws)
    header_row = [cell.value for cell in ws[header_row_number]]
    idx = {name: i for i, name in enumerate(header_row) if name in ("House", "Visit 2")}
    for row in ws.iter_rows(min_row=header_row_number + 1, values_only=True):
        if row[idx["House"]] == house_id:
            v2 = parse_visit_value(row[idx["Visit 2"]], house_id, "Visit 2")
            return v2 - PDT_OFFSET if v2 else None
    return None


def load_house_df(csv_file):
    df = pd.read_csv(csv_file, parse_dates=["timestamp_utc"])
    df["t"] = df["timestamp_utc"] - PDT_OFFSET  # local time (PDT), for display only
    return df


def series(df, instrument, variable, phase=None):
    mask = (df["instrument"] == instrument) & (df["variable"] == variable)
    if phase:
        mask &= df["phase"] == phase
    sub = df[mask].sort_values("t")
    if sub.empty:
        return None, None, None
    return sub["t"], sub["value"], sub["unit"].iloc[0]


def mark_stove_install(ax, stove_time):
    if stove_time is not None:
        ax.axvline(stove_time, color="red", linestyle="--", linewidth=1.5, alpha=0.85, zorder=0)


def label_pre_post(ax, df):
    pre = df[df["phase"] == "pre"]["t"]
    post = df[df["phase"] == "post"]["t"]
    if pre.empty or post.empty:
        return
    pre_center = pre.min() + (pre.max() - pre.min()) / 2
    post_center = post.min() + (post.max() - post.min()) / 2
    ax.text(pre_center, -0.15, "pre", transform=ax.get_xaxis_transform(),
            ha="center", va="top", color="dimgray", fontsize=9, clip_on=False)
    ax.text(post_center, -0.15, "post", transform=ax.get_xaxis_transform(),
            ha="center", va="top", color="dimgray", fontsize=9, clip_on=False)


def plot_panel(ax, df, instrument, variable, stove_time):
    x, y, unit = series(df, instrument, variable)
    if x is None:
        ax.set_title(f"{instrument}: {variable} -- NO DATA FOUND")
        return
    ax.plot(x, y, linewidth=0.8, color="tab:blue")
    ax.set_ylabel(f"{variable}\n({unit})" if unit else variable)
    ax.set_title(f"{instrument}: {variable}")
    ax.grid(alpha=0.3)
    mark_stove_install(ax, stove_time)


def plot_dual_panel(ax, df, left_spec, right_spec, stove_time):
    """left_spec/right_spec are (instrument, variable, color) tuples,
    plotted on the same panel with two separate y-axes."""
    li, lv, lcolor = left_spec
    x, y, unit = series(df, li, lv)
    if x is not None:
        ax.plot(x, y, linewidth=0.8, color=lcolor, label=f"{li}: {lv}")
        ax.set_ylabel(f"{lv}\n({unit})" if unit else lv, color=lcolor)
        ax.tick_params(axis="y", labelcolor=lcolor)
    else:
        print(f"Warning: no data for {li}:{lv}")

    ri, rv, rcolor = right_spec
    ax2 = ax.twinx()
    x2, y2, unit2 = series(df, ri, rv)
    if x2 is not None:
        ax2.plot(x2, y2, linewidth=0.8, color=rcolor, label=f"{ri}: {rv}")
        ax2.set_ylabel(f"{rv}\n({unit2})" if unit2 else rv, color=rcolor)
        ax2.tick_params(axis="y", labelcolor=rcolor)
    else:
        print(f"Warning: no data for {ri}:{rv}")

    ax.set_title(f"{li} {lv} + {ri} {rv}")
    ax.grid(alpha=0.3)
    mark_stove_install(ax, stove_time)


def plot_multi_left_one_right_panel(ax, df, left_specs, right_spec, stove_time):
    """left_specs is a list of (instrument, variable, color, label) all on
    the same left axis; right_spec is one (instrument, variable, color) on
    a secondary axis."""
    left_unit = None
    for instrument, variable, color, label in left_specs:
        x, y, unit = series(df, instrument, variable)
        if x is None:
            print(f"Warning: no data for {instrument}:{variable}")
            continue
        ax.plot(x, y, linewidth=0.8, color=color, label=label)
        left_unit = left_unit or unit
    ax.set_ylabel(f"stove temp\n({left_unit})" if left_unit else "stove temp")
    ax.legend(loc="upper left", fontsize=8)

    ri, rv, rcolor = right_spec
    ax2 = ax.twinx()
    x2, y2, unit2 = series(df, ri, rv)
    if x2 is not None:
        ax2.plot(x2, y2, linewidth=0.8, color=rcolor, label=f"{ri}: {rv}")
        ax2.set_ylabel(f"{rv}\n({unit2})" if unit2 else rv, color=rcolor)
        ax2.tick_params(axis="y", labelcolor=rcolor)
        ax2.set_ylim(0, 1500)  # fixed scale, per team request -- keeps this comparable across houses
    else:
        print(f"Warning: no data for {ri}:{rv}")

    ax.set_title(f"Geocene stove temps + {ri} {rv}")
    ax.grid(alpha=0.3)
    mark_stove_install(ax, stove_time)


def save_figure(fig, axes, df, house_id, title, output_path):
    label_pre_post(axes[-1], df)
    axes[-1].set_xlabel("Time (PDT)")
    fig.suptitle(f"House {house_id} -- {title}", fontsize=13, fontweight="bold")
    plt.tight_layout()
    plt.savefig(output_path, dpi=150)
    plt.close(fig)
    print(f"Saved {output_path}")


def make_airassure_figure(df, house_id, stove_time, output_dir):
    variables = ["temperature", "humidity", "co2", "pm25", "no2"]
    fig, axes = plt.subplots(len(variables), 1, figsize=(12, 3 * len(variables)), sharex=True)
    for ax, var in zip(axes, variables):
        plot_panel(ax, df, "AirAssure", var, stove_time)
    save_figure(fig, axes, df, house_id, "AirAssure", os.path.join(output_dir, f"H{house_id}_airassure.png"))


def make_atmocube_figure(df, house_id, stove_time, output_dir):
    variables = ["temperature", "humidity", "co2", "pm25", "nox_index"]
    fig, axes = plt.subplots(len(variables), 1, figsize=(12, 3 * len(variables)), sharex=True)
    for ax, var in zip(axes, variables):
        plot_panel(ax, df, "Atmocube", var, stove_time)
    save_figure(fig, axes, df, house_id, "Atmocube", os.path.join(output_dir, f"H{house_id}_atmocube.png"))


def make_aranet_kestrel_figure(df, house_id, stove_time, output_dir):
    fig, axes = plt.subplots(3, 1, figsize=(12, 9), sharex=True)
    plot_dual_panel(axes[0], df, ("Aranet", "temperature", "tab:blue"), ("Aranet", "humidity", "tab:orange"), stove_time)
    plot_dual_panel(axes[1], df, ("Kestrel", "temperature", "tab:blue"), ("Kestrel", "humidity", "tab:orange"), stove_time)
    plot_panel(axes[2], df, "Aranet", "co2", stove_time)
    save_figure(fig, axes, df, house_id, "Aranet + Kestrel", os.path.join(output_dir, f"H{house_id}_aranet_kestrel.png"))


def make_hobo_geocene_anemometer_figure(df, house_id, stove_time, output_dir):
    fig, axes = plt.subplots(2, 1, figsize=(12, 6), sharex=True)
    plot_multi_left_one_right_panel(
        axes[0], df,
        left_specs=[
            ("Geocene", "stove_temp_left", "tab:blue", "Geocene left"),
            ("Geocene", "stove_temp_right", "tab:cyan", "Geocene right"),
        ],
        right_spec=("Hobo", "power", "tab:red"),
        stove_time=stove_time,
    )
    plot_panel(axes[1], df, "Anemometer", "hood_airflow", stove_time)
    save_figure(fig, axes, df, house_id, "Hobo + Geocene + Anemometer", os.path.join(output_dir, f"H{house_id}_hobo_geocene_anemometer.png"))


def build_arg_parser():
    parser = argparse.ArgumentParser(description="Generate the 4 standardized QA/QC figures for one house")
    parser.add_argument("csv_file", help="Path to the H<N>.csv file")
    parser.add_argument("house_charac_file", help="Path to the House Characteristics .xlsx file (for the stove-install time)")
    parser.add_argument("--output-dir", default="qc_plots", help="Folder to save the 4 figures into")
    return parser


def main():
    args = build_arg_parser().parse_args()
    os.makedirs(args.output_dir, exist_ok=True)

    df = load_house_df(args.csv_file)
    house_id = df["house_id"].iloc[0]
    house_num = house_id.lstrip("H")
    stove_time = get_stove_install_time(args.house_charac_file, house_id)
    if stove_time is None:
        print(f"Warning: couldn't find Visit 2 for {house_id} in {args.house_charac_file} -- no red line will be drawn")

    make_airassure_figure(df, house_num, stove_time, args.output_dir)
    make_atmocube_figure(df, house_num, stove_time, args.output_dir)
    make_aranet_kestrel_figure(df, house_num, stove_time, args.output_dir)
    make_hobo_geocene_anemometer_figure(df, house_num, stove_time, args.output_dir)


if __name__ == "__main__":
    main()